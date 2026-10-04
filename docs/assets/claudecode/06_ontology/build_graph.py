"""엑셀 실적 데이터에서 지식 그래프 개체·관계를 뽑아 ontology.html에 넣는다.

사용법:
    python3 build_graph.py          # 검증만 (개체·관계 수 출력)
    python3 build_graph.py --write  # ontology.html의 GRAPH_DATA 블록 교체

규칙은 ontology.html 06 규칙(R1~R9)을 따른다.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
XLSX = HERE.parent / "06_dashboard" / "데이터" / "피자 매장별 실적 데이터.xlsx"
HTML = HERE / "ontology.html"

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)


def sheet(name):
    rows = list(wb[name].iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:] if r[0]]


stores, menus, sups, pfs, events = (sheet(n) for n in ["매장마스터", "메뉴마스터", "공급사마스터", "플랫폼마스터", "사건로그"])
sales = list(wb["매출데이터"].iter_rows(values_only=True))
H = {k: i for i, k in enumerate(sales[0])}
sales = sales[1:]

nodes, edges = [], []


def node(id_, label, type_, **kw):
    nodes.append({"id": id_, "label": label, "type": type_, **kw})


def edge(s, t, rel, kind="direct"):
    edges.append({"s": s, "t": t, "rel": rel, "kind": kind})


# R1 매장 동일성: 코드·구코드·매장명·별칭 → ST 코드
alias = {}
for s in stores:
    code = s["매장코드"]
    keys = [code, s["구매장코드(2022-08 이전)"], s["매장명"]] + [a.strip() for a in s["별칭(다른 시스템 표기)"].split(";")]
    for k in keys:
        alias[k] = code
SCOPE = {"본사", "전 매장"}


def resolve(text):
    """표기 문자열 → (매장코드 집합, 범위 여부)"""
    codes, scope = set(), False
    for part in str(text or "").split(";"):
        p = part.strip()
        if p in SCOPE:
            scope = True
        elif p in alias:
            codes.add(alias[p])
    return codes, scope


# 조직
node("ALL", "전 매장", "hub", sub="본사·전 매장 범위 (9곳)")
for hq in dict.fromkeys(s["지역본부"] for s in stores):
    node(f"HQ:{hq}", hq, "hq")
for sv in dict.fromkeys((s["담당SV코드"], s["담당SV"]) for s in stores):
    node(sv[0], sv[1], "sv", sub=sv[0])
for s in stores:
    c = s["매장코드"]
    node(c, s["매장명"], "store", sub=f"{c} · {s['가맹유형']} · {s['상권유형']}")
    edge(c, f"HQ:{s['지역본부']}", "소속")
    edge(c, s["담당SV코드"], "담당 SV")
    edge("ALL", c, "포함")
    owner = f"OW:{c}"
    since = s["점주 변경일"]
    node(owner, s["점주/점장"], "owner", sub=f"{s['매장명']} 점주" + (f" ({since}~)" if since else ""))
    edge(c, owner, "운영 점주")
    if since:
        prev = f"OW:{c}:1"
        node(prev, "1대 점주", "owner", sub=f"{s['매장명']} ~{since} · 이름 기록 없음")
        edge(c, prev, "운영 점주 (종료)")

# 상품·거래
CATS = ["피자", "사이드메뉴", "세트메뉴", "음료", "기타(굿즈/쿠폰)"]
for c in CATS:
    node(f"CAT:{c}", c, "cat")
    edge(f"CAT:{c}", "ALL", "매출 (카테고리매출 경유)", "derived")
for p in pfs:
    node(p["플랫폼코드"], p["플랫폼명"], "pf", sub=f"{p['플랫폼코드']} · 수수료 {p['현재 중개수수료율']:.1%}")
for s in sups:
    node(s["공급사코드"], s["공급사명"], "sup", sub=f"{s['공급사코드']} · {s['품목']}")
for m in menus:
    code = m["메뉴코드"]
    if code.startswith("FEE"):
        node(code, m["메뉴명"], "fee", sub="요금항목 · 상품 아님 (R2)")
        continue
    end = f" · {m['단종일']} 단종" if m["단종일"] else ""
    node(code, m["메뉴명"], "menu", sub=f"{code} · {m['카테고리']} · {m['판매 주기']}{end}", cat=m["카테고리"])
    edge(code, f"CAT:{m['카테고리']}", "분류")
    for sup in re.findall(r"SUP-\d+", m["주요 식자재 공급사(코드)"] or ""):
        edge(code, sup, "식자재 공급")
    for part in (m["세트 구성 메뉴(코드)"] or "").split(";"):
        alts = re.findall(r"MN-\d+", part)
        for a in alts:
            edge(code, a, "구성 (택1)" if len(alts) > 1 else "구성")

# 사건
GROUP = {
    "운영": "매장 운영", "위생": "매장 운영", "클레임": "매장 운영", "인력": "매장 운영", "리뉴얼": "매장 운영", "점주": "매장 운영",
    "외부환경": "외부 요인", "경쟁": "외부 요인", "플랫폼": "외부 요인", "공급사": "외부 요인",
    "본사정책": "본사 결정", "가격정책": "본사 결정", "시스템": "본사 결정", "메뉴": "본사 결정", "인력정책": "본사 결정", "채널": "본사 결정",
    "신메뉴": "상품", "데이터": "데이터",
}
for e in events:
    eid = e["사건ID"]
    conf = "확인" if e["확인여부"].startswith("확인") else ("미확인" if e["확인여부"].startswith("미확인") else "추정")
    node(eid, e["제목"], "event", date=e["일자"].strftime("%Y-%m-%d"), group=GROUP[e["유형"]], conf=conf,
         sub=f"{eid} · {e['일자']:%Y-%m-%d} · {e['유형']} · {conf}")
    codes, scope = resolve(e["매장(기록 표기)"])
    c2, s2 = resolve(e["관련 대상"])
    codes |= c2
    for c in sorted(codes):
        edge(eid, c, "발생 매장")
    if (scope or s2) and not codes:
        edge(eid, "ALL", "발생 범위")
    for t in re.findall(r"(?:MN|SUP|PF|FEE)-\d+", e["관련 대상"] or ""):
        edge(eid, t, "관련 대상")

# 사건 간 관계 (L16, 사건로그 내용에서 읽어 낸 추론 관계)
for a, b, rel in [
    ("EV-004", "EV-009", "선행"),   # 대구 시범 결과를 근거로 전국 확대
    ("EV-009", "EV-037", "선행"),   # 2021년부터 매년 여름 한정
    ("EV-006", "EV-007", "선행"),   # 베타 → 정식 출시
    ("EV-022", "EV-023", "선행"),   # 리뉴얼 후 클레임 최저
    ("EV-025", "EV-035", "같은 유형"),  # 연동점 정규직 이탈 반복
    ("EV-032", "EV-033", "선행"),   # 트러플 출시 → 홍대점 수요 폭증
    ("EV-033", "EV-034", "선행"),   # 주방 과부하 → 배달 지연 클레임
    ("EV-016", "EV-013", "같은 유형"),  # "강남점 2022년 사례와 같은 유형"
    ("EV-034", "EV-016", "같은 유형"),  # 홍대점 배달 지연 클레임 반복
]:
    edge(a, b, rel, "inferred")
# 공급사마스터 근거 사건 (L13과 양쪽 일치 확인)
for s in sups:
    if s["근거 사건"]:
        assert any(e["s"] == s["근거 사건"] and e["t"] == s["공급사코드"] for e in edges), s

# 질문 답의 근거가 되는 월간실적 개체 (R3: 매장·월 1건)
# Q1: 2024-10 대비 플랫폼수수료율(÷매출합)이 오르고 영업이익률이 내린 2024-11~12 매장·월
sm = {}
for r in sales:
    sm[(r[H["매장코드"]], r[H["날짜(YYYY-MM)"]])] = r
EVIDENCE = [
    ("ST-005", "2024-11", ["EV-024"]), ("ST-005", "2024-12", ["EV-024"]), ("ST-006", "2024-11", ["EV-024"]), ("ST-007", "2024-12", ["EV-024"]),
    ("ST-002", "2026-03", ["EV-033"]), ("ST-002", "2026-04", ["EV-034"]),
]
name = {s["매장코드"]: s["매장명"] for s in stores}
for code, ym, evs in EVIDENCE:
    r = sm[(code, ym)]
    assert r[H["구분"]] == "실적"
    sid = f"SM:{code}:{ym}"
    node(sid, f"{name[code].split()[1]} {ym}", "sm",
         sub=f"이익률 {r[H['영업이익률(%)']]:.1f}% · 플랫폼수수료 {r[H['플랫폼수수료(만원)']]:,.0f}만원 · 클레임 {r[H['클레임건수']]:.0f}건")
    edge(sid, code, "매장")
    for ev in evs:
        edge(ev, sid, "영향 기간")
    if "EV-024" in evs:
        edge(sid, "PF-01", "주문 경유")

ids = {n["id"] for n in nodes}
bad = [e for e in edges if e["s"] not in ids or e["t"] not in ids]
assert not bad, bad
counts = {}
for n in nodes:
    counts[n["type"]] = counts.get(n["type"], 0) + 1
print(f"nodes {len(nodes)} {counts}")
print(f"edges {len(edges)}")

if "--write" in sys.argv:
    data = json.dumps({"nodes": nodes, "edges": edges}, ensure_ascii=False, separators=(",", ":"))
    html = HTML.read_text(encoding="utf-8")
    html, n = re.subn(r"(/\*GRAPH_DATA_START\*/).*?(/\*GRAPH_DATA_END\*/)", lambda m: m.group(1) + data + m.group(2), html, flags=re.S)
    assert n == 1, "GRAPH_DATA 블록을 찾지 못함"
    HTML.write_text(html, encoding="utf-8")
    print("ontology.html 갱신")
