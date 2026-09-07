#!/usr/bin/env python3
"""
매출데이터 시트를 지정 월까지 연장한다.

마지막 월 다음 달부터 END_YM까지, 매장마다 전년 같은 달 실적에
그 매장의 최근 성장률(최근 5개월 / 전년 같은 5개월)을 곱해 만든다.
카테고리 6행 구조, 매장·월 열의 6행 반복, 정수·소수 자릿수는 2025년 행 형식을 따른다.

이야기 연속성(사건로그와 맞춤):
- 2026-06 슈림프 골드 여름 한정 재판매 → 전 매장 신메뉴 비중 12% 안팎
- 2026-07 폭염 → 배달 비중 +4%p, 홀 −4%p
- 2026-08 제주 관광 성수기 → 연동점 매출 +10%
- 2026-06~ 광주 충장로점 2대 점주 체제 계속 → 인건비 +6%, 영업이익률 −3%p, QSC −4

사용법: python3 extend_sales_data.py           # 미리보기(추가될 행 수와 요약)
        python3 extend_sales_data.py --write   # 시트에 추가
"""
import sys, random
from pathlib import Path
import pandas as pd
import openpyxl
from copy import copy

HERE = Path(__file__).parent
XLSX = HERE / '데이터' / '피자 매장별 실적 데이터.xlsx'
SHEET = '매출데이터'
END_YM = '2026-08'
CATS = ['기타(굿즈/쿠폰)', '배달수수료', '사이드메뉴', '세트메뉴', '음료', '피자']
COST_COLS = ['식자재비(만원)', '인건비(만원)', '임차료(만원)', '플랫폼수수료(만원)', '로열티(만원)', '광고분담금(만원)']
random.seed(20260907)


def noise(pct):
    return 1 + random.uniform(-pct, pct)


def main(write):
    df = pd.read_excel(XLSX, sheet_name=SHEET)
    df['ym'] = df['날짜(YYYY-MM)'].astype(str).str[:7]
    last = df.ym.max()
    months = pd.period_range(pd.Period(last, 'M') + 1, pd.Period(END_YM, 'M'), freq='M')
    if len(months) == 0:
        print(f'이미 {last}까지 있습니다.'); return
    sm = df.drop_duplicates(['매장명', 'ym']).set_index(['매장명', 'ym'])
    rev = df.groupby(['매장명', 'ym'])['매출액(만원)'].sum()
    cat = df.groupby(['매장명', 'ym', '매출카테고리'])['매출액(만원)'].sum()
    stores = sorted(df.매장명.unique())
    cols = [c for c in df.columns if c != 'ym']
    new_rows = []
    summary = []
    for st in stores:
        recent = [f'2026-{m:02d}' for m in range(1, 6)]
        prev = [f'2025-{m:02d}' for m in range(1, 6)]
        growth = rev.loc[st, recent].mean() / rev.loc[st, prev].mean()
        for p in months:
            ym = str(p); base_ym = f'{p.year - 1}-{p.month:02d}'
            b = sm.loc[(st, base_ym)]
            cur = sm.loc[(st, '2026-05')]
            story_rev = 1.0
            if ym == '2026-08' and st == '제주 연동점': story_rev = 1.10
            cat_rev = {}
            for c in CATS:
                cat_rev[c] = int(round(cat.loc[(st, base_ym, c)] * growth * story_rev * noise(0.03)))
            total = sum(cat_rev.values())
            orders = int(round(b['총주문건수'] * growth * story_rev * noise(0.02)))
            hall, deli, take = b['홀_주문비중'], b['배달_주문비중'], b['포장_주문비중']
            if ym == '2026-07': hall -= 0.04; deli += 0.04
            hall, deli = round(hall, 2), round(deli, 2); take = round(1 - hall - deli, 3)
            labor_mult = 1.06 if st == '광주 충장로점' else 1.0
            costs = {
                '식자재비(만원)': int(round(b['식자재비(만원)'] * growth * story_rev * noise(0.03))),
                '인건비(만원)': int(round(b['인건비(만원)'] * growth * labor_mult * noise(0.03))),
                '임차료(만원)': int(round(cur['임차료(만원)'])),
                '플랫폼수수료(만원)': int(round(b['플랫폼수수료(만원)'] * growth * (1.15 if b.name[1] < '2024-11' else 1.0) * noise(0.03))),
                '로열티(만원)': int(round(total * 0.055)),
                '광고분담금(만원)': int(round(total * 0.02)),
            }
            # 영업이익률은 전년 같은 달 수준에서 출발한다(원본도 매출−비용과 정확히 일치하지 않음).
            margin = b['영업이익률(%)'] + (-3.0 if st == '광주 충장로점' else 0.0) + random.uniform(-1.5, 1.5)
            margin = round(margin, 1)
            profit = int(round(total * margin / 100))
            newmenu = round(0.12 * noise(0.15), 3) if ym == '2026-06' else round(min(b['신메뉴매출비중'], 0.08) * noise(0.2), 3)
            qsc = round(b['QSC점수'] + (-4 if st == '광주 충장로점' else 0) + random.uniform(-1.5, 1.5), 1)
            base_row = {
                '매장명': st, '연도': 2026, '월': p.month, '날짜(YYYY-MM)': ym,
                '영업일수': int(b['영업일수']), '총주문건수': orders, '객단가(만원)': round(total / orders, 1),
                '홀_주문비중': hall, '배달_주문비중': deli, '포장_주문비중': take,
                '배민_주문수': int(round(b['배민_주문수'] * growth * noise(0.03))),
                '쿠팡이츠_주문수': int(round(b['쿠팡이츠_주문수'] * growth * noise(0.03))),
                '요기요_주문수': int(round(b['요기요_주문수'] * growth * noise(0.03))),
                '점심피크_매출비중': b['점심피크_매출비중'], '저녁피크_매출비중': b['저녁피크_매출비중'], '야간_매출비중': b['야간_매출비중'],
                **costs,
                '영업이익(만원)': profit, '영업이익률(%)': margin,
                '신규고객수': int(round(b['신규고객수'] * growth * noise(0.05))),
                '재방문율': round(min(0.7, cur['재방문율'] * noise(0.02)), 3),
                '쿠폰할인금액(만원)': int(round(b['쿠폰할인금액(만원)'] * growth * 1.05 * noise(0.05))),
                '자사앱주문비중': round(min(0.4, cur['자사앱주문비중'] * noise(0.03)), 3),
                '리뷰평점': round(b['리뷰평점'] + random.uniform(-0.1, 0.1), 1),
                '식자재발주금액(만원)': int(round(costs['식자재비(만원)'] * 1.03 * noise(0.02))),
                '식자재폐기율': round(b['식자재폐기율'] * noise(0.2), 3),
                '신메뉴매출비중': newmenu,
                'QSC점수': qsc, '위생검사결과': b['위생검사결과'],
                '클레임건수': int(round(b['클레임건수'] * noise(0.2))),
                '정규직원수': int(round(cur['정규직원수'])), '파트타임직원수': int(round(cur['파트타임직원수'])),
            }
            for c in CATS:
                row = dict(base_row); row['매출카테고리'] = c; row['매출액(만원)'] = cat_rev[c]
                new_rows.append([row[k] for k in cols])
            summary.append((st, ym, total, margin, hall, deli, newmenu))
    print(f'현재 마지막 월 {last} → {END_YM}, 추가 {len(new_rows)}행')
    print(pd.DataFrame(summary, columns=['매장', '월', '매출합', '이익률', '홀', '배달', '신메뉴']).to_string(index=False))
    if not write:
        print('미리보기만 했습니다. --write 로 시트에 추가합니다.'); return
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    # 서식은 2025년 행(텍스트 날짜·정수 값)에서 복사
    ref = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 4).value) == '2025-06': ref = r; break
    for vals in new_rows:
        ws.append(vals)
        r = ws.max_row
        for c in range(1, len(vals) + 1):
            src = ws.cell(ref, c); dst = ws.cell(r, c)
            dst.font = copy(src.font); dst.number_format = src.number_format
            dst.alignment = copy(src.alignment); dst.border = copy(src.border); dst.fill = copy(src.fill)
    wb.save(XLSX)
    print(f'저장: {SHEET} {ws.max_row - 1}행')


if __name__ == '__main__':
    main('--write' in sys.argv)
